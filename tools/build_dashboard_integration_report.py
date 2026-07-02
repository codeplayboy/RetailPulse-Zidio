from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\rasha\Downloads\RetailPulse\RetailPulse-Zidio")
OUT = ROOT / "docs" / "RetailPulse_Dashboard_Integration_Testing_Report_Rashad.docx"
EVIDENCE = ROOT / "docs" / "dashboard_evidence"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(width_dxa))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_run(run, bold=False, italic=False, size=None, color=None):
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")


def add_para(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        style_run(r, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        style_run(r2)
    else:
        r = p.add_run(text)
        style_run(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        style_run(r)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        style_run(r)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    if widths is None:
        widths = [9360 // len(headers)] * len(headers)

    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        set_cell_width(hdr[i], widths[i])
        set_cell_margins(hdr[i])
        set_cell_shading(hdr[i], "F2F4F7")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(text))
        style_run(r, bold=True, size=9.5, color="1F4D78")

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_width(cells[i], widths[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 or len(str(value)) > 24 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            style_run(r, size=9)
    doc.add_paragraph()
    return table


def add_callout(doc, title, body, fill="F4F6F9"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_width(cell, 9360)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    style_run(r, bold=True, color="1F4D78")
    p2 = cell.add_paragraph()
    r2 = p2.add_run(body)
    style_run(r2)
    doc.add_paragraph()


def add_figure(doc, image_name, caption):
    path = EVIDENCE / image_name
    if not path.exists():
        add_para(doc, f"Screenshot missing: {image_name}")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.25))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    style_run(r, italic=True, size=9, color="555555")


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("RetailPulse Dashboard Integration Report - Prepared by Rashad")
    style_run(r, size=9, color="555555")


def build_report():
    doc = Document()
    configure_doc(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run("RetailPulse Dashboard Integration, Testing, and Evidence Report")
    style_run(r, bold=True, size=24, color="0B2545")

    subtitle = doc.add_paragraph()
    r = subtitle.add_run("Prepared by Rashad - Platform Integration / Dashboard")
    style_run(r, size=12, color="555555")
    r2 = subtitle.add_run(f"\nGenerated: {datetime.now().strftime('%B %d, %Y, %I:%M %p')}")
    style_run(r2, size=10, color="555555")

    add_callout(
        doc,
        "Executive summary",
        "I verified the RetailPulse dashboard as the integration layer for the team project. The application builds successfully, serves locally with HTTP 200, loads the dashboard public data files, renders the main business modules, and plays distinct Remotion data stories for Churn Prediction and Demand Forecasting. The current remaining issue is not a dashboard integration blocker: the local Python environment does not have pytest installed, so the included Python test files compile but the pytest runner cannot execute until the environment is completed.",
        fill="E8EEF5",
    )

    doc.add_heading("1. Project Context", level=1)
    add_para(doc, "RetailPulse is an AI-powered customer analytics and demand forecasting platform. My role is the platform integration and dashboard layer: I connect the data engineering, customer intelligence, demand forecasting, churn, inventory, reporting, and media-story outputs into a single executive dashboard.")
    add_table(
        doc,
        ["Team member", "Project role", "Integrated dashboard responsibility"],
        [
            ["Rashad", "Platform Integration / Dashboard", "React dashboard, visual system, module integration, live data binding, Remotion Play Data stories, testing evidence, final documentation."],
            ["Kaviya", "Data Engineering", "Revenue, demand, SKU inventory forecasting datasets, cleaning summary, enriched customer inventory data, and dashboard-ready processed CSV sources."],
            ["Rohinee", "Data Pipeline / Customer Intelligence", "Cleaned data, feature engineering, RFM/customer features, full features_data.csv shared separately due GitHub size limit, GitHub-friendly features_data_sample.csv."],
            ["Sachin", "Demand Forecasting / ML Outputs", "Prophet forecasting pipeline, demand_forecast.csv, inventory_recommendations.csv, churn_predictions.csv, high_risk_customers.csv, model evaluation and limitations."],
        ],
        widths=[1500, 2100, 5760],
    )

    doc.add_heading("2. Integration Architecture", level=1)
    add_para(doc, "The dashboard is built as a React + TypeScript + Vite application under dashboard/. It does not depend on fake hardcoded dashboard totals as the primary source. The live dashboard data loader reads CSV and JSON files from dashboard/public/data/ and transforms them into typed dashboard state.")
    add_para(doc, "Current frontend stack:")
    add_bullets(doc, [
        "React 19 and TypeScript for the application shell and module views.",
        "Vite for local development, production build, and preview serving.",
        "GSAP and Framer Motion for entrance motion, hover states, scroll animation, and premium UI transitions.",
        "Three.js / React Three Fiber for 3D-style visual elements.",
        "Remotion Player for in-dashboard Play Data story videos.",
        "Lucide React for consistent professional interface icons.",
    ])
    add_para(doc, "Data flow implemented in the dashboard:")
    add_numbered(doc, [
        "Data files are placed in dashboard/public/data/ so the Vite app can fetch them at runtime.",
        "dashboard/src/data/retailpulse-data.ts loads revenue, demand, inventory, forecast, churn, high-risk customer, feature sample, and feature summary sources.",
        "The loader parses CSV/JSON, aggregates KPIs, builds chart series, prepares report previews, and provides module-specific rows.",
        "dashboard/src/App.tsx renders the navigation, top bar, hero, module pages, filters, KPI cards, charts, tables, media studio, and settings.",
        "dashboard/src/chart-animations.tsx powers Remotion data-story modals and module-level cinematic stories.",
    ])

    add_table(
        doc,
        ["Dashboard source", "Purpose", "Evidence"],
        [
            ["dashboard/public/data/dataset_1_daily_revenue_forecasting.csv", "Executive revenue, monthly performance, KPI totals, trend charts.", "Loaded by retailpulse-data.ts fetch layer."],
            ["dashboard/public/data/dataset_2_daily_demand_forecasting.csv", "Demand forecasting context, weekly/rolling demand signals.", "Loaded by retailpulse-data.ts fetch layer."],
            ["dashboard/public/data/dataset_3_sku_level_inventory_forecasting.csv", "SKU/product and inventory trend context.", "Loaded by retailpulse-data.ts fetch layer."],
            ["dashboard/public/data/demand_forecast.csv", "Forecast horizon rows, confidence bands, forecast table, media story data.", "147 rendered rows in dashboard public copy."],
            ["dashboard/public/data/churn_predictions.csv", "Churn KPIs, risk segmentation, retention rows, churn page.", "5000 customer rows in dashboard public copy."],
            ["dashboard/public/data/inventory_recommendations.csv", "Inventory health, reorder recommendations, stock alerts, inventory table.", "4917 recommendation rows in dashboard public copy."],
            ["dashboard/public/data/features_data_sample.csv", "Dashboard-safe schema sample for feature engineering integration tests.", "10000 sample rows with full feature schema."],
            ["dashboard/public/data/features_summary.json", "Summary of full feature dataset for dashboard-scale KPIs.", "JSON summary available and loaded optionally."],
        ],
        widths=[3000, 3600, 2760],
    )

    doc.add_heading("3. Team Module Integration Status", level=1)
    add_table(
        doc,
        ["Module", "Owner", "Integrated artifacts", "Dashboard status"],
        [
            ["Data Engineering", "Kaviya", "Daily revenue, daily demand, SKU inventory forecasting datasets, enriched customer inventory support.", "Integrated through dashboard/public/data CSV copies and transformed into revenue, demand, category, and SKU views."],
            ["Data Pipeline and Features", "Rohinee", "cleaned_data.csv, features_data.csv, features_data_sample.csv, features_summary.json, cleaning and feature engineering source files.", "Full features_data.csv is local outside GitHub because it is around 188 MB. Dashboard uses sample plus summary for safe browser integration."],
            ["Demand Forecasting", "Sachin", "demand_forecast.csv, Prophet production code, metrics MAE 5534.49, RMSE 8759.64, MAPE 22.53%.", "Integrated into Demand Forecasting page, forecast table, hero KPIs, Remotion forecast story, and Media Studio demand story."],
            ["Inventory Optimization", "Sachin", "inventory_recommendations.csv with safety stock, reorder point, inventory gap, status, and recommendation fields.", "Integrated into Inventory Optimization page, gauge score, stock alerts, recommendations table, and report outputs."],
            ["Churn Prediction", "Sachin plus customer data pipeline", "churn_predictions.csv and high_risk_customers.csv with customer churn/risk fields.", "Integrated into Churn Prediction page, churn risk KPIs, risk analysis, retention recommendations, and unique Remotion story videos."],
            ["Platform Integration", "Rashad", "React dashboard, data loader, chart components, Remotion Play Data videos, settings, reports, media studio, responsive UI.", "Verified by build, browser smoke tests, screenshots, and module walkthrough."],
        ],
        widths=[1750, 1300, 3650, 2660],
    )

    doc.add_heading("4. Interface Walkthrough", level=1)
    add_para(doc, "The dashboard is organized as an enterprise command center with a sidebar, top search/command bar, hero section, replay controls, KPI cards, charts, tables, data stories, and module-specific analytics panels.")
    add_table(
        doc,
        ["Section", "What it does", "Outcome"],
        [
            ["Executive Overview", "Shows revenue, orders, customer health, churn risk, revenue trend, category mix, product ranking, activity feed, monthly performance, and 3D volume analysis.", "Gives a board-ready view of the business state."],
            ["Customer Segmentation", "Uses RFM/customer intelligence concepts, segment tabs, CLV ranking, cluster/force visualizations, territory maps, and playbook cards.", "Turns Rohinee's feature engineering/customer intelligence work into operational customer strategy."],
            ["Churn Prediction", "Shows churn probability distribution, risk matrix, feature importance, churn timeline, neural network visualization, and retention recommendations.", "Connects churn outputs to clear retention actions."],
            ["Demand Forecasting", "Shows Prophet confidence, projected demand, MAPE context, scenario controls, historical demand plus forecast, seasonal and weekly trend panels, forecast detail, category forecast.", "Connects Sachin's Prophet work to planning decisions."],
            ["Inventory Optimization", "Shows inventory health gauge, stock alerts, SKU tracking, critical/understock/overstock cards, reorder planning, and recommendations table.", "Converts demand/inventory outputs into replenishment decisions."],
            ["Analytics & Reports", "Creates customer, sales, churn, forecast, and inventory report previews from dashboard data.", "Supports export-ready reporting and final presentation needs."],
            ["Media Studio", "Provides Remotion module videos for executive storytelling.", "Supports presentation-style data films for customer, churn, demand, inventory, and report modules."],
            ["Settings", "Contains theme toggle, model settings, forecast settings, user/team management, and notification controls.", "Keeps controls in a professional settings area rather than cluttering the main UI."],
        ],
        widths=[1800, 4600, 2960],
    )

    doc.add_heading("5. Visual Evidence", level=1)
    add_para(doc, "The following screenshots were captured from the live local preview at http://127.0.0.1:4173/ using Chrome/Playwright and the in-app browser. Browser console error checks returned zero errors during the capture run.")
    add_figure(doc, "desktop_01_overview.png", "Figure 1. Executive Overview loaded from integrated dashboard data.")
    add_figure(doc, "desktop_03_churn.png", "Figure 2. Churn Prediction page with live risk KPIs and retention command visuals.")
    add_figure(doc, "desktop_04_churn_story.png", "Figure 3. Churn Play Data story using the unique churn probability distribution Remotion visual.")
    add_figure(doc, "desktop_05_forecasting.png", "Figure 4. Demand Forecasting page showing Prophet confidence, MAPE status, projected demand, and scenario controls.")
    add_figure(doc, "desktop_06_forecast_story.png", "Figure 5. Demand Play Data story using the unique forecast horizon Remotion visual.")
    add_figure(doc, "desktop_07_inventory.png", "Figure 6. Inventory Optimization page with stock health and SKU recommendations.")
    add_figure(doc, "desktop_08_reports.png", "Figure 7. Analytics and Reports module for export-ready report packs.")
    add_figure(doc, "desktop_09_media.png", "Figure 8. Media Studio module for Remotion story generation.")
    add_figure(doc, "desktop_10_settings.png", "Figure 9. Settings page containing model controls, team access, theme controls, and notification controls.")
    add_figure(doc, "mobile_01_overview.png", "Figure 10. Mobile responsive overview after data load.")

    doc.add_heading("6. Testing Performed", level=1)
    add_table(
        doc,
        ["Test", "Command / method", "Result", "Interpretation"],
        [
            ["Production build", "npm run build in dashboard/", "Passed: TypeScript build and Vite production build completed. 2002 modules transformed.", "The React/TypeScript dashboard compiles successfully."],
            ["Local preview endpoint", "Invoke-WebRequest http://127.0.0.1:4173/", "Passed: HTTP 200 and RetailPulse title found.", "The built preview is reachable locally."],
            ["Browser smoke test", "Chrome/Playwright page walkthrough across Overview, Segmentation, Churn, Forecasting, Inventory, Reports, Media Studio, Settings.", "Passed: all target pages opened and screenshots captured.", "The main interface routes and command search flow work."],
            ["Browser console", "Chrome/Playwright console error collection during screenshot run.", "Passed: zero console errors reported.", "No visible runtime JavaScript errors during the tested walkthrough."],
            ["Remotion Churn story", "Opened Play Data Story: Churn Probability Distribution.", "Passed: unique churn probability ring video rendered.", "The repeated generic story bug was fixed for churn."],
            ["Remotion Forecast story", "Opened Play Data Story: Historical Demand + Forecast.", "Passed: unique forecast horizon video rendered.", "The repeated generic story bug was fixed for demand forecasting."],
            ["Python syntax check", "py -3 -m compileall src tests", "Passed: source modules and test files compiled.", "Python files are syntactically valid."],
            ["Pytest execution", "py -3 -m pytest tests -q", "Blocked: No module named pytest.", "Automated tests exist, but the local Python environment needs pytest installed before running them."],
            ["Dashboard data file verification", "Line/header checks on CSV and JSON files.", "Passed for present files.", "Dashboard has the required public data files and module output sources."],
        ],
        widths=[1800, 3000, 2500, 2060],
    )

    doc.add_heading("7. Data Verification Summary", level=1)
    add_table(
        doc,
        ["File", "Rows / lines verified", "Schema evidence"],
        [
            ["data/processed/cleaned_data.csv", "779,424 data rows plus header", "Invoice, StockCode, Description, Quantity, InvoiceDate, Price, Customer_ID, Country, TotalPrice"],
            ["C:/Users/rasha/Downloads/features_data.csv", "779,424 data rows plus header; about 188 MB", "RFM, customer, lag, rolling, purchase, and profit feature columns present."],
            ["data/processed/features_data_sample.csv", "10,000 sample rows plus header", "Same feature schema as full file for GitHub-safe testing."],
            ["data/outputs/demand_forecast.csv", "739 data rows plus header", "ds, yhat, yhat_lower, yhat_upper"],
            ["data/outputs/churn_predictions.csv", "5,000 data rows plus header", "Customer_ID, behavior fields, Churn_Risk, Churn, Risk_Level"],
            ["data/outputs/inventory_recommendations.csv", "4,917 data rows plus header", "Safety_Stock, Reorder_Point, Inventory_Gap, Inventory_Status, Recommendation"],
            ["dashboard/public/data/demand_forecast.csv", "147 data rows plus header", "Date, Predicted_Demand, Lower_Bound, Upper_Bound"],
            ["dashboard/public/data/churn_predictions.csv", "5,000 data rows plus header", "Dashboard-ready churn customer records."],
            ["dashboard/public/data/inventory_recommendations.csv", "4,917 data rows plus header", "Dashboard-ready inventory recommendation records."],
            ["dashboard/public/data/features_summary.json", "48 lines", "Feature summary keys available for dashboard-scale integration."],
        ],
        widths=[3500, 2500, 3360],
    )

    doc.add_heading("8. Challenges Faced and Fixes Applied", level=1)
    add_table(
        doc,
        ["Challenge", "Cause", "Fix / current handling"],
        [
            ["Large feature dataset could not be pushed to GitHub.", "features_data.csv is about 188 MB, above GitHub's 100 MB normal file limit.", "Rohinee shared the full file separately. A GitHub-friendly features_data_sample.csv and features_summary.json are used for repository and dashboard integration testing."],
            ["Dashboard initially used placeholder/fake-looking values.", "The UI layer was ahead of the actual teammate output integration.", "Created a data loader that reads public CSV/JSON outputs and derives KPIs, charts, tables, report rows, and media story inputs."],
            ["Repeated Remotion Play Data videos.", "Multiple dashboard panels were routed to generic line/network/timeline story templates.", "Added unique Churn and Demand story visuals and explicit story IDs for each panel."],
            ["Churn story first-frame labels looked incorrect.", "Risk labels were multiplied by the reveal animation progress.", "Kept ring animation but made labels display normalized final values from frame start."],
            ["Header/content scroll overlap appeared during earlier UI testing.", "Sticky/fixed topbar and scroll container spacing were conflicting.", "Adjusted layout behavior so content starts below the header and scrolls through a single clean workspace flow."],
            ["Background animation felt distracting.", "Animated background effects competed with dashboard content.", "Replaced moving background with low-poly SVG backgrounds for dark and light themes."],
            ["Dashboard performance concerns from heavy motion.", "Too many expensive filters/background animations can cause jank.", "Reduced background animation load, used GPU-friendly transform/opacity motion, kept production build chunks separated."],
            ["Pytest could not run locally.", "Both active Python interpreters lack pytest; default python points to a Layla/Hermes virtual environment.", "Recorded blocker. Python compileall still verifies syntax. Install pytest in the RetailPulse environment before final automated Python test run."],
            ["LibreOffice was missing for document render QA.", "soffice was not on PATH initially.", "Installed LibreOffice through winget; soffice.exe is now present at C:/Program Files/LibreOffice/program/soffice.exe."],
        ],
        widths=[2800, 3100, 3460],
    )

    doc.add_heading("9. Current Readiness Assessment", level=1)
    add_callout(
        doc,
        "My assessment",
        "The dashboard integration is functionally ready for presentation and team review. The frontend builds, the local preview works, the key module pages render, the data files are connected, and the Remotion story repetition issue has been fixed. Before final submission, I would still complete one final clean-environment run with pytest installed and ask each teammate to confirm their final CSV schemas remain unchanged.",
        fill="F2F4F7",
    )
    add_table(
        doc,
        ["Area", "Completion", "Notes"],
        [
            ["Dashboard UI and navigation", "95%", "All core pages exist and render; polish pass complete enough for review."],
            ["Data integration", "92%", "Main CSV/JSON outputs are wired. Full features data is external due file size; summary/sample are dashboard-safe."],
            ["Demand forecasting integration", "95%", "Forecast output, metrics, limitations, page UI, and Play Data story integrated."],
            ["Churn integration", "95%", "Churn predictions, risk breakdown, retention actions, and unique Remotion stories integrated."],
            ["Inventory integration", "95%", "Inventory recommendations and health logic are visible in dashboard."],
            ["Testing evidence", "90%", "Build/browser/source checks done; pytest runner still needs environment setup."],
            ["Documentation", "95%", "This report documents architecture, integration, testing, screenshots, challenges, and outcomes."],
        ],
        widths=[3000, 1500, 4860],
    )

    doc.add_heading("10. Final Next Steps", level=1)
    add_numbered(doc, [
        "Install pytest and required Python dependencies inside the correct RetailPulse environment, then run py -3 -m pytest tests -q again.",
        "Ask teammates to confirm no final schema changes are pending for demand_forecast.csv, churn_predictions.csv, inventory_recommendations.csv, cleaned_data.csv, and features_data_sample.csv.",
        "Commit the dashboard integration, documentation, and evidence files from feature/rashad-platform-engineering.",
        "Deploy the latest build to Vercel or GitHub Pages and run one final deployed-browser smoke test.",
        "Use this report as the dashboard documentation section for the final project submission or PDF.",
    ])

    doc.add_heading("Appendix A. Key Local Evidence Paths", level=1)
    add_table(
        doc,
        ["Evidence", "Path"],
        [
            ["Dashboard source", str(ROOT / "dashboard" / "src")],
            ["Dashboard public data", str(ROOT / "dashboard" / "public" / "data")],
            ["Python modules", str(ROOT / "src")],
            ["Python tests", str(ROOT / "tests")],
            ["Evidence screenshots", str(EVIDENCE)],
            ["Full external feature file", r"C:\Users\rasha\Downloads\features_data.csv"],
            ["This document", str(OUT)],
        ],
        widths=[2600, 6760],
    )

    doc.add_heading("Appendix B. Commands Used for Verification", level=1)
    commands = [
        "npm run build",
        "Invoke-WebRequest -Uri http://127.0.0.1:4173/ -UseBasicParsing",
        "py -3 -m compileall src tests",
        "py -3 -m pytest tests -q",
        "Chrome/Playwright screenshot walkthrough at http://127.0.0.1:4173/",
    ]
    add_bullets(doc, commands)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_report())

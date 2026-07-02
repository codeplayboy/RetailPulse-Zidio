from __future__ import annotations

from datetime import datetime
from pathlib import Path
import sys

sys.path.append(str(Path(__file__).parent))

from build_dashboard_integration_report import (  # noqa: E402
    ROOT,
    EVIDENCE,
    add_bullets,
    add_callout,
    add_figure,
    add_numbered,
    add_para,
    add_table,
    configure_doc,
    style_run,
)

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


OUT = ROOT / "docs" / "RetailPulse_Complete_Dashboard_Integration_Book_Rashad_Roushan.docx"


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("RetailPulse Complete Dashboard Integration Book")
    style_run(r, bold=True, size=25, color="0B2545")
    p2 = doc.add_paragraph()
    r2 = p2.add_run("Architecture, module integration, testing proof, screenshots, challenges, fixes, and final readiness")
    style_run(r2, size=12, color="555555")
    p3 = doc.add_paragraph()
    r3 = p3.add_run(f"Prepared by: Rashad Roushan | Platform Integration and Dashboard\nGenerated: {datetime.now().strftime('%B %d, %Y, %I:%M %p')}")
    style_run(r3, size=10, color="555555")


def build_book() -> Path:
    doc = Document()
    configure_doc(doc)
    add_title(doc)

    add_callout(
        doc,
        "Purpose of this book",
        "I prepared this document as the complete proof book for my RetailPulse dashboard work. It explains how the dashboard was built, how each teammate's module output is connected, how the interface works, what was tested, what proof was captured, what problems we faced, and what fixes were applied. It is written from my perspective as Rashad Roushan, responsible for platform integration and dashboard delivery.",
        fill="E8EEF5",
    )

    doc.add_heading("1. Executive Summary", level=1)
    add_para(
        doc,
        "RetailPulse is now functioning as an integrated retail intelligence dashboard. The React/Vite dashboard reads module output files from dashboard/public/data, transforms them into analytics state, and renders the project modules as an executive interface. The dashboard includes customer intelligence, churn prediction, demand forecasting, inventory optimization, analytics/reporting, media stories, and settings.",
    )
    add_para(
        doc,
        "The strongest proof is not just that the UI opens. The dashboard-specific data contract check passes for every public data file used by the application, the production frontend build passes, the local preview returns HTTP 200, browser walkthrough screenshots were captured for all major pages, Remotion Play Data stories were verified for Churn Prediction and Demand Forecasting, and the Python test suite now passes after schema/date normalization.",
    )
    add_table(
        doc,
        ["Proof area", "Result", "Meaning"],
        [
            ["Dashboard production build", "Passed", "TypeScript and Vite compile the dashboard successfully."],
            ["Live preview", "Passed", "http://127.0.0.1:4173/ responds with HTTP 200 and loads RetailPulse."],
            ["Dashboard data contract", "Passed", "Revenue, demand, SKU, forecast, churn, inventory, and feature files required by the dashboard are present with required columns."],
            ["Browser walkthrough", "Passed", "Overview, Segmentation, Churn, Forecasting, Inventory, Reports, Media Studio, Settings, and mobile layout screenshots captured."],
            ["Remotion stories", "Passed", "Churn and Forecast Play Data stories open and render unique visuals."],
            ["Python source compile", "Passed", "src/ and tests/ Python files compile."],
            ["Pytest suite", "Passed", "20 tests passed after fixing schema references and normalizing Excel serial dates."],
        ],
        widths=[2800, 1800, 4760],
    )

    doc.add_heading("2. My Role and Team Integration", level=1)
    add_para(doc, "My role is Platform Integration and Dashboard. I am responsible for bringing together the work from the data pipeline, machine learning outputs, inventory recommendations, churn predictions, and forecasting outputs into one working dashboard that can be presented and reviewed.")
    add_table(
        doc,
        ["Team member", "Contribution", "How I integrated it"],
        [
            ["Rashad Roushan", "Dashboard, integration layer, visual system, Play Data stories, documentation, testing evidence.", "Built the React/Vite interface, data loader, module views, Remotion stories, settings, reports, and evidence book."],
            ["Kaviya", "Data engineering datasets: daily revenue, daily demand, SKU-level inventory forecasting, cleaning summary, customer inventory enrichment.", "Used dashboard/public/data revenue, demand, and SKU CSVs to drive overview, demand, inventory, product, and chart sections."],
            ["Rohinee", "Data cleaning, feature engineering, cleaned_data.csv, features_data.csv, features_data_sample.csv, feature summary.", "Used the cleaned data and feature summary/sample approach for customer intelligence and dashboard-safe integration because full features_data.csv is too large for normal GitHub push."],
            ["Sachin", "Prophet demand forecasting, churn predictions, inventory recommendations, model metrics and limitations.", "Integrated demand_forecast.csv, churn_predictions.csv, high_risk_customers.csv, and inventory_recommendations.csv into Forecasting, Churn, Inventory, Reports, and Media Studio."],
        ],
        widths=[1800, 3700, 3860],
    )

    doc.add_heading("3. System Architecture From Scratch", level=1)
    add_numbered(
        doc,
        [
            "Raw and processed data are produced by teammate modules under data/processed and data/outputs.",
            "Dashboard-safe copies are stored under dashboard/public/data so the browser can fetch them at runtime.",
            "dashboard/src/data/retailpulse-data.ts is the integration adapter. It fetches CSV/JSON files, parses them, validates available data, derives KPIs, builds chart arrays, and returns typed DashboardData.",
            "dashboard/src/App.tsx is the interface shell. It renders navigation, search, hero sections, module pages, cards, filters, tables, settings, and command modal.",
            "dashboard/src/chart-animations.tsx powers Remotion Player videos for panel-level Play Data stories and module-level media stories.",
            "dashboard/src/App.css and component files define the final glassmorphism/dark luxury visual system, responsive layout, and motion states.",
        ],
    )
    add_figure(doc, "proof_source_binding.png", "Proof A. Source binding evidence showing runtime CSV/JSON fetches from the dashboard integration adapter.")

    doc.add_heading("4. Module-by-Module Integration Evidence", level=1)
    add_para(doc, "This section documents exactly what each dashboard module proves and which teammate outputs it depends on.")
    add_table(
        doc,
        ["Dashboard module", "Integrated files", "Rendered evidence"],
        [
            ["Executive Overview", "dataset_1_daily_revenue_forecasting.csv, churn_predictions.csv, inventory_recommendations.csv", "Revenue, orders, customers, churn index, revenue trend, product/category analysis, activity feed."],
            ["Customer Segmentation", "features_data_sample.csv, features_summary.json, churn/customer fields", "RFM/segment controls, CLV ranking, customer playbook, cluster-style views."],
            ["Churn Prediction", "churn_predictions.csv, high_risk_customers.csv", "Risk distribution, risk matrix, feature importance, churn timeline, neural model, retention recommendations."],
            ["Demand Forecasting", "demand_forecast.csv, dataset_2_daily_demand_forecasting.csv", "Projected demand, forecast confidence, MAPE card, scenario controls, historical + forecast chart, forecast table."],
            ["Inventory Optimization", "inventory_recommendations.csv, SKU inventory dataset", "Inventory health, critical stock, understock, overstock, stock alerts, reorder plan."],
            ["Analytics & Reports", "All integrated datasets", "Customer, sales, churn, forecast, and inventory report previews."],
            ["Media Studio", "DashboardData transformed into Remotion props", "Video stories for customer segments, churn intelligence, demand forecast, inventory command, and reports."],
            ["Settings", "App state and team/project configuration", "Theme toggle, model settings, forecast settings, user management, notification settings."],
        ],
        widths=[2200, 3800, 3360],
    )

    doc.add_heading("5. Screenshots of Working Interfaces", level=1)
    add_para(doc, "These screenshots were captured from the live dashboard preview, not mocked separately. They show the integrated modules rendering in the actual app.")
    screenshot_items = [
        ("desktop_01_overview.png", "Executive Overview with revenue, orders, customer, and churn metrics loaded."),
        ("desktop_02_segmentation.png", "Customer Segmentation interface showing customer intelligence module."),
        ("desktop_03_churn.png", "Churn Prediction interface with risk metrics and retention command visuals."),
        ("desktop_04_churn_story.png", "Churn Play Data story proving unique Remotion story integration."),
        ("desktop_05_forecasting.png", "Demand Forecasting page with Prophet/MAPE context and forecast controls."),
        ("desktop_06_forecast_story.png", "Demand Forecast Play Data story proving unique forecast animation integration."),
        ("desktop_07_inventory.png", "Inventory Optimization page showing inventory health and stock planning."),
        ("desktop_08_reports.png", "Analytics and Reports module with export/report-pack interface."),
        ("desktop_09_media.png", "Media Studio module for Remotion-based video stories."),
        ("desktop_10_settings.png", "Settings interface with model, forecast, user, and notification controls."),
        ("mobile_01_overview.png", "Responsive mobile overview proof after data load."),
    ]
    for name, caption in screenshot_items:
        add_figure(doc, name, caption)

    doc.add_heading("6. Proof of Integrated Data Files", level=1)
    add_para(doc, "The dashboard data contract check validates the actual files consumed by the browser application. This is separate from teammate Python tests because the dashboard uses standardized dashboard-facing column names.")
    add_figure(doc, "proof_dashboard_data_contract.png", "Proof B. Dashboard data contract results: all public dashboard data sources passed required column checks.")
    add_figure(doc, "proof_dataset_receipt.png", "Proof C. Dataset receipt proof including full features_data.csv size and dashboard-safe sample/summary handling.")
    add_table(
        doc,
        ["Public dashboard file", "Rows verified", "Purpose"],
        [
            ["dataset_1_daily_revenue_forecasting.csv", "739", "Revenue, orders, KPI and revenue trend calculations."],
            ["dataset_2_daily_demand_forecasting.csv", "739", "Demand trend and demand forecasting context."],
            ["dataset_3_sku_level_inventory_forecasting.csv", "4917", "SKU/product/inventory analytics source."],
            ["demand_forecast.csv", "147", "Forecast rows and demand story input."],
            ["churn_predictions.csv", "5000", "Churn risk, customer segments, retention recommendations."],
            ["inventory_recommendations.csv", "4917", "Stock health, reorder points, safety stock, inventory gaps."],
            ["features_data_sample.csv", "10000", "Feature schema sample for browser-safe integration testing."],
            ["features_summary.json", "Present", "Feature summary for dashboard-scale customer intelligence."],
        ],
        widths=[3600, 1600, 4160],
    )

    doc.add_heading("7. Testing Evidence and Results", level=1)
    add_figure(doc, "proof_npm_build.png", "Proof D. Production build output from npm run build.")
    add_figure(doc, "proof_pytest_findings.png", "Proof E. Pytest output showing real test findings after installing required Python packages.")
    add_table(
        doc,
        ["Test area", "Command", "Outcome"],
        [
            ["Frontend production build", "npm run build", "Passed. Vite transformed 2002 modules and emitted production assets."],
            ["Local dashboard preview", "Invoke-WebRequest http://127.0.0.1:4173/", "Passed. HTTP 200 and RetailPulse title found."],
            ["Browser walkthrough", "Chrome/Playwright module navigation and screenshot capture", "Passed. Screenshots captured for all main modules and no console errors reported."],
            ["Dashboard data contract", "py -3 tools/dashboard_data_contract_check.py", "Passed. All dashboard-facing CSV sources and required columns were present."],
            ["Python syntax", "py -3 -m compileall src tests", "Passed. Python source and test files compiled."],
            ["Python test suite", "py -3 -m pytest tests -q", "Passed after fixes. Result: 20 passed."],
        ],
        widths=[2500, 3100, 3760],
    )

    doc.add_heading("8. Problems, Struggles, and Fixes", level=1)
    add_para(doc, "This project had real integration issues. I am documenting them clearly because they show how the final dashboard became stable.")
    add_table(
        doc,
        ["Problem faced", "Impact", "How it was fixed or handled"],
        [
            ["Full features_data.csv was too large for GitHub.", "Rohinee's full feature output could not be pushed normally because it is about 188 MB.", "Kept full file locally/shared separately; used features_data_sample.csv and features_summary.json for repository-safe and dashboard-safe integration."],
            ["Column naming mismatch between tests and generated files.", "Original tests expected Customer ID, Lag1, Lag7, while generated/current files use Customer_ID, Lag_1, Lag_7.", "Updated tests to use the standardized project schema."],
            ["features_data.csv not present under data/processed in repo.", "The full feature file is around 188 MB and should not be committed normally.", "Updated tests to use the full file when present, otherwise validate data/processed/features_data_sample.csv. Dashboard uses sample + summary for browser-safe integration."],
            ["Python environment initially lacked pandas, numpy, pytest.", "Tests could not run at all.", "Installed the required Python packages, reran pytest, and verified 20 passing tests."],
            ["InvoiceDate fields were converted incorrectly.", "The cleaned/features files contained Excel serial day values interpreted as nanoseconds after 1970, creating 1970 dates.", "Recovered the Excel serial day values and normalized repo-safe files into real 2009-2011 dates; feature date columns now match."],
            ["Default python path earlier resolved to an unrelated environment.", "Commands could accidentally use the wrong Python stack.", "Verified system Python with py -3 and documented the environment issue."],
            ["Repeated Remotion Play Data videos.", "Churn and Demand panels looked like copied story templates.", "Added explicit story IDs and unique visuals for churn distribution, risk matrix, retention timeline, forecast horizon, seasonality, weekly cycle, and category forecast."],
            ["Charts initially looked plain/incomplete.", "Some charts lacked endpoint markers, overflow control, or professional animation.", "Improved chart styling, bloom-on-scroll animation, point markers, and responsive chart containment."],
            ["Gauge needle was detached earlier.", "Inventory gauge looked broken and visually inaccurate.", "Rebuilt gauge as a cohesive component with shared center coordinates for arc/needle."],
            ["Header overlapped content while scrolling.", "Cards scrolled underneath the header.", "Fixed scroll container spacing and stacking behavior so content starts below the header."],
            ["Background animations were distracting.", "Moving spirals and heavy background effects competed with dashboard content.", "Replaced them with low-poly SVG backgrounds for dark/light themes."],
            ["Light theme palette was too dark.", "Light mode looked muddy and misaligned with reference visuals.", "Adjusted light theme toward brighter metallic/silver styling while preserving contrast."],
            ["LibreOffice headless render hung locally.", "DOCX-to-PNG render QA could not complete in this Windows environment.", "Performed structural DOCX QA and embedded evidence images; documented render limitation transparently."],
        ],
        widths=[3000, 3000, 3360],
    )

    doc.add_heading("9. Interface Explanation in Detail", level=1)
    add_para(doc, "The interface is intentionally built as a live command center rather than a basic admin dashboard.")
    add_table(
        doc,
        ["Interface element", "Purpose", "Why it matters"],
        [
            ["Sidebar navigation", "Switches between Overview, Segmentation, Churn, Forecasting, Inventory, Reports, Media Studio, and Settings.", "Keeps all project modules visible and demonstrates that my dashboard is the integration layer."],
            ["Top search/command bar", "Provides module search and quick command access.", "Makes the dashboard feel like an operating system and helps navigate during presentation."],
            ["Hero module panel", "Summarizes each module with primary, operational, and trend state metrics.", "Gives a fast executive read before detailed charts."],
            ["KPI cards", "Show major business metrics with readable trend indicators.", "Makes module outcomes easy to explain to evaluators."],
            ["Charts and tables", "Render revenue, demand, categories, churn, stock, reports, and recommendations.", "Turns CSV outputs into visible insight."],
            ["Play Data buttons", "Open Remotion-powered story videos for charts/panels.", "Provides proof that dashboard data can be transformed into presentation-ready data films."],
            ["Media Studio", "Centralizes module videos and export-style story formats.", "Useful for final demo and project presentation."],
            ["Settings", "Holds model, forecast, team, notification, and theme controls.", "Keeps controls professional and avoids clutter on analytics pages."],
        ],
        widths=[2300, 3600, 3460],
    )

    doc.add_heading("10. Current Final Assessment", level=1)
    add_callout(
        doc,
        "Final assessment by Rashad Roushan",
        "The dashboard is integrated and ready for presentation/review: live dashboard files are connected, the UI builds, the browser walkthrough passes, modules render with proof screenshots, the dashboard data contract passes, and the Python test suite now passes. The only remaining team decision is where to store the full 188 MB features_data.csv for final submission.",
        fill="F4F6F9",
    )
    add_table(
        doc,
        ["Area", "Readiness", "Reason"],
        [
            ["Dashboard UI", "Ready", "All pages render and screenshots prove working state."],
            ["Frontend build", "Ready", "Production build passes."],
            ["Dashboard data integration", "Ready", "Dashboard data contract check passes."],
            ["Remotion Play Data", "Ready", "Unique Churn and Demand stories verified."],
            ["Team module outputs", "Ready", "Main output files exist and are integrated; dashboard contract passes."],
            ["Python tests", "Ready", "20 tests passed after schema/date fixes."],
            ["Documentation", "Ready", "This book contains architecture, screenshots, proof, test outputs, challenges, fixes, and next steps."],
        ],
        widths=[2600, 1600, 5160],
    )

    doc.add_heading("11. Final Next Steps Before Submission", level=1)
    add_numbered(
        doc,
        [
            "Decide how to store full features_data.csv: Git LFS, external Drive/OneDrive link, or keep only sample + summary in GitHub.",
            "Ask teammates to confirm the standardized schema names stay stable: Customer_ID, Lag_1, Lag_7, RollingMean7, RollingStd7.",
            "Keep the date-normalization helper documented so the Excel serial issue does not return.",
            "Run pytest again after any teammate output changes and document the final passing result.",
            "Deploy the latest dashboard build and run one final public URL browser smoke test.",
            "Use this book as the dashboard/integration documentation artifact for the final project report.",
        ],
    )

    doc.add_heading("Appendix A. Evidence File Locations", level=1)
    add_table(
        doc,
        ["Evidence type", "Location"],
        [
            ["Complete book", str(OUT)],
            ["Evidence screenshots", str(EVIDENCE)],
            ["Dashboard data contract JSON", str(EVIDENCE / "dashboard_data_contract_results.json")],
            ["NPM build output", str(EVIDENCE / "npm_build_result.txt")],
            ["Pytest output", str(EVIDENCE / "pytest_result.txt")],
            ["Dashboard integration adapter", str(ROOT / "dashboard/src/data/retailpulse-data.ts")],
            ["React app shell", str(ROOT / "dashboard/src/App.tsx")],
            ["Remotion stories", str(ROOT / "dashboard/src/chart-animations.tsx")],
            ["Full external features data", r"C:\Users\rasha\Downloads\features_data.csv"],
        ],
        widths=[3000, 6360],
    )

    doc.add_heading("Appendix B. Commands Run", level=1)
    add_bullets(
        doc,
        [
            "npm run build",
            "Invoke-WebRequest -Uri http://127.0.0.1:4173/ -UseBasicParsing",
            "py -3 tools/dashboard_data_contract_check.py",
            "py -3 -m compileall src tests",
            "py -3 -m pytest tests -q",
            "Chrome/Playwright screenshot walkthrough across dashboard modules",
        ],
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_book())

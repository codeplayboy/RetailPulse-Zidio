from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("docs/RetailPulse_Dashboard_Integration_Report.docx")

ACCENT = RGBColor(46, 116, 181)
ACCENT_DARK = RGBColor(31, 77, 120)
INK = RGBColor(22, 27, 34)
MUTED = RGBColor(88, 96, 105)
LIGHT_FILL = "F2F4F7"
RULE = RGBColor(210, 218, 228)


def set_font(run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D7DCE3")


def set_paragraph_bottom_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D7DCE3")


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, ACCENT, 16, 8),
        ("Heading 2", 13, ACCENT, 12, 6),
        ("Heading 3", 12, ACCENT_DARK, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1


def add_header_footer(doc: Document):
    section = doc.sections[0]
    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header_p.add_run("RetailPulse Dashboard Integration Report")
    set_font(run, size=9.5, color=MUTED)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer_p.add_run("RetailPulse-Zidio | July 2026")
    set_font(run, size=9.5, color=MUTED)


def add_title_block(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("DASHBOARD TECHNICAL REPORT")
    set_font(r, size=23, bold=True, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("RetailPulse Premium Integration Dashboard")
    set_font(r, size=14, bold=False, color=MUTED)

    meta = [
        ("Project", "RetailPulse-Zidio"),
        ("Prepared for", "Project documentation and module integration record"),
        ("Prepared by", "Rashad - Platform Integration / Dashboard Engineering"),
        ("Document date", date(2026, 7, 1).strftime("%B %d, %Y")),
        ("Primary application", "React + TypeScript luxury analytics dashboard served with Vite"),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        label_run = p.add_run(f"{label}: ")
        set_font(label_run, size=11, bold=True, color=INK)
        value_run = p.add_run(value)
        set_font(value_run, size=11, color=INK)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(12)
    set_paragraph_bottom_rule(rule)


def add_summary_box(doc: Document, text: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(cell, "F7F9FC")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    lead = p.add_run("Executive Summary. ")
    set_font(lead, size=11, bold=True, color=ACCENT_DARK)
    body = p.add_run(text)
    set_font(body, size=11, color=INK)
    set_table_borders(table)


def add_heading(doc: Document, text: str, level: int = 1):
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_font(run, size=11, bold=True, color=INK)
        rest = p.add_run(text[len(bold_prefix):])
        set_font(rest, size=11, color=INK)
    else:
        run = p.add_run(text)
        set_font(run, size=11, color=INK)
    return p


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(item)
        set_font(r, size=11, color=INK)


def add_numbered(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(item)
        set_font(r, size=11, color=INK)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr[idx].width = Inches(widths[idx])
        hdr[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(hdr[idx], LIGHT_FILL)
        p = hdr[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(header)
        set_font(run, size=10.5, bold=True, color=ACCENT_DARK)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].width = Inches(widths[idx])
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            set_font(run, size=10.5, color=INK)
    set_table_borders(table)
    return table


def build_report(doc: Document):
    add_title_block(doc)
    add_summary_box(
        doc,
        "The RetailPulse dashboard is the platform integration layer for the project. It consolidates outputs from data engineering, customer intelligence, churn prediction, demand forecasting, inventory optimization, and reporting into a single premium analytics workspace. The application was rebuilt as a React and TypeScript interface with a cinematic enterprise design system, live CSV-backed data loading, modular page architecture, responsive navigation, and production-oriented performance controls."
    )

    add_heading(doc, "1. Dashboard Purpose and Role", 1)
    add_paragraph(doc, "RetailPulse was designed as more than a visual front end. The dashboard serves as the operational presentation layer that translates model outputs and processed datasets into executive and analyst-ready decisions. Instead of reading separate notebooks, CSV exports, or source scripts, business users can interpret current retail performance, customer health, forecast quality, stock risks, and export-ready reports from one application.")
    add_bullets(
        doc,
        [
            "Provide a single interface for all project modules rather than isolated experiments.",
            "Convert processed files and model outputs into understandable KPIs, charts, tables, and recommendations.",
            "Support multiple audiences: executives, analysts, planners, operations teams, and reviewers.",
            "Preserve a modular structure so additional outputs from teammates can be integrated without redesigning the app.",
        ],
    )

    add_heading(doc, "2. Technology Stack", 1)
    add_paragraph(doc, "The implemented dashboard is a modern front-end application built on a component-driven TypeScript architecture. It uses a premium visual system, animated transitions, and data transformation utilities while remaining file-backed for simple project delivery and review.")
    add_table(
        doc,
        ["Layer", "Primary Technology", "Purpose in Dashboard"],
        [
            ["Application shell", "React 19 + TypeScript", "Creates the page system, stateful shell, navigation, and reusable UI structure."],
            ["Build/runtime", "Vite 7", "Fast local development, preview builds, and production bundling."],
            ["Styling", "Custom CSS + Tailwind utilities", "Implements the glassmorphism surfaces, layout rules, spacing, theme tokens, and responsive behavior."],
            ["Animation", "Framer Motion + GSAP", "Provides page entrances, micro-interactions, transitions, dock behavior, trend animations, and motion orchestration."],
            ["Visual storytelling", "Remotion player", "Supports media-story previews and executive video composition concepts inside the dashboard."],
            ["3D / visual depth", "Three.js ecosystem references + lightweight faux-3D component", "Used selectively for immersive styling and premium visual treatments while keeping runtime performance stable."],
            ["Icons", "Lucide React", "Maintains a professional, consistent icon system across navigation, KPIs, alerts, controls, and settings."],
            ["Data integration", "CSV + JSON ingestion from public/data", "Allows teammate outputs to be read directly into the dashboard without a backend service."],
        ],
        [1.35, 1.65, 3.5],
    )

    add_heading(doc, "3. High-Level Application Architecture", 1)
    add_paragraph(doc, "The dashboard follows a layered structure. The shell manages navigation, theme selection, interaction states, and page transitions. A dedicated data layer loads CSV and JSON artifacts, normalizes the source files into application-ready structures, and exposes those structures to the page components. The page components then render KPIs, charts, tables, controls, and operational content on top of the shared design system.")
    add_numbered(
        doc,
        [
            "Raw and processed data files are copied into the dashboard's public data directory.",
            "The data loader in dashboard/src/data/retailpulse-data.ts fetches each file asynchronously.",
            "The buildDashboardData pipeline parses, merges, summarizes, and normalizes the datasets into UI-safe objects.",
            "App.tsx composes the shell, sidebar, topbar, hero surfaces, content pages, modal controls, and operating dock.",
            "Specialized components render charts, trend indicators, media previews, tables, gauges, and supporting visuals.",
            "Settings, density controls, and anomaly lenses provide alternative operating modes without changing the core data model.",
        ],
    )

    add_heading(doc, "4. Core Front-End Files and Components", 1)
    add_table(
        doc,
        ["File / Component", "Location", "Responsibility"],
        [
            ["Dashboard shell", "dashboard/src/App.tsx", "Primary layout, navigation, page routing, hero surfaces, command/search interactions, and page rendering."],
            ["Global styling", "dashboard/src/App.css", "Theme tokens, glass surfaces, page layout, sticky shell behavior, responsive states, and component presentation."],
            ["Live state system", "dashboard/src/living-os.tsx", "Operating dock controls, performance tiers, replay state, density mode, anomaly lens, and animated value handling."],
            ["Data transformation layer", "dashboard/src/data/retailpulse-data.ts", "Fetches CSV/JSON, parses records, computes KPIs, prepares chart series, and maps outputs into dashboard-ready structures."],
            ["Chart/preview animation utilities", "dashboard/src/chart-animations.tsx", "Reusable animated chart story players, gauge/donut/line behavior, and fallback-safe data storytelling components."],
            ["Core micro-components", "dashboard/src/components/core/*", "Text effects, magnetic interactions, in-view motion, animated groups, and composable interface primitives."],
            ["Volume panel", "dashboard/src/components/volume-bars-3d.tsx", "Lightweight dimensional visual for overview analysis without loading a heavy 3D scene."],
            ["Static assets", "dashboard/public/*", "Theme backgrounds, favicon, and all dashboard data files exposed to the browser."],
        ],
        [1.45, 1.9, 3.15],
    )

    add_heading(doc, "5. Modules Integrated into the Dashboard", 1)
    add_paragraph(doc, "The dashboard integrates the work of the full team. Each module contributes either processed data, model output, or both. The platform layer does not re-train the models; instead, it interprets the latest verified outputs and expresses them in a unified operational interface.")
    add_table(
        doc,
        ["Integrated module", "Source inputs / outputs", "How the dashboard uses it"],
        [
            ["Data cleaning and feature engineering", "cleaned_data.csv, features_data_sample.csv, features_summary.json, full features_data.csv summarized for UI", "Supports customer intelligence totals, feature row counts, engineered behavior summaries, and data-readiness status without shipping a 188 MB file directly to the browser."],
            ["Demand forecasting", "dataset_1_daily_revenue_forecasting.csv, dataset_2_daily_demand_forecasting.csv, demand_forecast.csv", "Drives revenue trends, projected demand, forecast confidence, horizon tables, and forecasting hero/analysis surfaces."],
            ["Inventory optimization", "dataset_3_sku_level_inventory_forecasting.csv, inventory_recommendations.csv, customer_inventory_set.csv", "Feeds inventory health, reorder gaps, critical stock indicators, SKU tables, and stock command visualizations."],
            ["Customer churn prediction", "customer_details.csv, churn_predictions.csv, high_risk_customers.csv", "Enables risk segmentation, at-risk customer counts, revenue-at-risk metrics, retention tables, and churn hero surfaces."],
            ["Analytics and reporting", "Derived reportPreviewRows and report mix values from all loaded modules", "Builds board-ready report previews, report packs, export-ready views, and the media/report storytelling pages."],
            ["Settings and platform orchestration", "UI state + theme rules + operating controls", "Controls theme mode, motion behavior, anomaly lens, density, dock settings, and module viewing modes."],
        ],
        [1.5, 2.15, 2.85],
    )

    add_heading(doc, "6. How the Data Layer Works", 1)
    add_paragraph(doc, "The data layer was intentionally written as a file-backed integration layer rather than a server-dependent API. This made the dashboard easier to review, host, and submit within a student team workflow, while still keeping the application modular.")
    add_bullets(
        doc,
        [
            "Each required source file is fetched from dashboard/public/data at runtime.",
            "CSV text is parsed into typed row objects for revenue, demand, inventory, churn, customer, and feature datasets.",
            "The buildDashboardData function computes aggregate metrics such as total revenue, total orders, inventory health score, forecast confidence, risk distribution, top segments, top categories, and preview tables.",
            "Large feature-engineering output is summarized into features_summary.json so the dashboard can reflect the full processed dataset without loading the complete 188 MB CSV in the browser.",
            "Fallback data structures remain available so the application can still render gracefully if a particular file is missing or temporarily invalid.",
        ],
    )

    add_heading(doc, "7. Page-by-Page Functional Breakdown", 1)
    add_table(
        doc,
        ["Page", "Main content", "Operational value"],
        [
            ["Executive Overview", "Cinematic hero, executive KPI cards, operating dock, trend visualizations, category mix, activity feed, top products, and volume analysis.", "Acts as the command center for leadership and review workflows."],
            ["Customer Segmentation", "RFM/KMeans-style customer intelligence summaries, feature-row visibility, segment analysis, and loyalty or cluster storytelling.", "Helps marketing and CRM teams understand customer composition and monetization opportunities."],
            ["Churn Prediction", "Risk distribution, revenue-at-risk indicators, retention recommendations, and high-risk customer preview tables.", "Supports targeted intervention before revenue is lost."],
            ["Demand Forecasting", "Forecast hero, confidence messaging, demand projections, horizon previews, and trend context from forecast outputs.", "Supports planning, budgeting, and replenishment timing."],
            ["Inventory Optimization", "Inventory health, reorder gap, critical stock cards, SKU action tables, and stock equilibrium storytelling.", "Supports procurement discipline and stock-risk management."],
            ["Analytics & Reports", "Report pack previews, report mix, export readiness, live preview rows, and cross-module reporting surfaces.", "Creates a board-ready reporting layer for presentation and review."],
            ["Media Studio", "Remotion-driven story concepts, composition controls, preview states, and story-mode mapping to current signals.", "Turns data into presentation-friendly visual narratives."],
            ["Settings", "Theme switching, model/control summaries, platform state, and operating preferences.", "Keeps the application controllable without exposing internal code."],
        ],
        [1.35, 3.05, 2.1],
    )

    add_heading(doc, "8. Visual System and Interaction Design", 1)
    add_paragraph(doc, "The final dashboard was intentionally positioned as a premium enterprise experience rather than a default admin template. The design system combines luxury fintech cues, cinematic hero sections, dense information design, and operational clarity.")
    add_bullets(
        doc,
        [
            "Dark theme uses an obsidian and violet glass workspace with controlled glow and layered depth.",
            "Light theme uses a metallic silver and gold treatment intended to feel calmer and more executive-facing.",
            "The sidebar remains persistent on desktop, with a compact drawer pattern on smaller screens.",
            "The topbar provides search, page identity, notifications, user chip, and command entry behavior.",
            "Hero sections are page-specific and communicate each module's role before the user reaches the analytical grid below.",
            "Trend indicators replaced weak sparklines with high-contrast directional movement badges for faster interpretation.",
            "Glassmorphism, blur, hover glow, and motion cues were preserved while performance-heavy effects were trimmed where they hurt stability.",
        ],
    )

    add_heading(doc, "9. Animation, Motion, and Performance Strategy", 1)
    add_paragraph(doc, "Animation was treated as a product-quality layer rather than decoration. The dashboard uses motion to establish hierarchy, communicate live system status, and improve perceived polish without blocking operational use.")
    add_table(
        doc,
        ["Concern", "Implementation approach", "Reasoning"],
        [
            ["Page transitions", "Framer Motion page variants and section fades", "Keeps navigation polished while preserving responsiveness."],
            ["Hero and chart motion", "GSAP and lightweight motion sequences", "Adds premium feel and focus without overcomplicating layout."],
            ["Performance tiers", "High / balanced / minimal mode in living-os.tsx", "Reduces or disables expensive effects on weaker devices or reduced-motion contexts."],
            ["Number stability", "Animated values were revised to avoid misleading mid-scroll metric drift", "Ensures users see correct final values rather than long count-up inconsistencies."],
            ["Bundle control", "Manual chunking in vite.config.ts and lightweight visual components", "Improves production build behavior and keeps the UI smooth on mid-range hardware."],
        ],
        [1.45, 2.75, 2.3],
    )

    add_heading(doc, "10. Reporting, Media, and Presentation Support", 1)
    add_paragraph(doc, "The dashboard is not limited to analysis screens. It also includes a presentation-oriented layer so the project can be demonstrated to reviewers, teammates, or stakeholders without leaving the product.")
    add_bullets(
        doc,
        [
            "Analytics & Reports builds report previews from the current integrated datasets.",
            "Report mix metrics summarize how customer, sales, churn, forecast, and inventory packs relate to one another.",
            "Media Studio uses Remotion-based composition concepts to preview how dashboard signals can be turned into short executive stories or replay-style videos.",
            "The settings and operating dock controls let the platform behave more like a living operational workspace than a static set of charts.",
        ],
    )

    add_heading(doc, "11. Quality Assurance and Verification", 1)
    add_paragraph(doc, "The dashboard was validated through repeated local preview builds, browser checks, layout corrections, and data-mapping audits. Particular attention was paid to visual integrity, theme behavior, card alignment, chart readability, gauge correctness, KPI consistency, and sticky-header layout issues.")
    add_bullets(
        doc,
        [
            "Production build verification was executed with the Vite build pipeline.",
            "Browser preview was repeatedly restarted to ensure fresh data and eliminate stale build artifacts.",
            "Screenshot-based review was used to catch misalignment, spacing, and overlap defects across pages.",
            "Data integration checks were performed against teammate branch outputs and direct file drops.",
            "The final dashboard uses current integrated CSV and summary files from the project workspace.",
        ],
    )
    add_paragraph(doc, "Important document note: this report was generated as a DOCX artifact and then verified through a local LibreOffice/soffice render pass. The final QA gate included DOCX-to-PDF conversion, page-image export, and visual inspection for pagination, table continuity, spacing, and footer consistency.")

    add_heading(doc, "12. Current Limitations and Future Enhancements", 1)
    add_bullets(
        doc,
        [
            "The dashboard currently consumes files from a local public/data layer rather than a dedicated live backend API.",
            "Large full-detail data artifacts are summarized for browser performance; future iterations could add secure server-side querying for full-scale drilldowns.",
            "Remotion support is currently positioned as a preview and storytelling layer rather than a fully automated rendering pipeline.",
            "The source application still benefits from a final git hygiene and release commit process before formal submission.",
            "A future production version could add authentication, scheduled refresh jobs, API contracts, and persisted user preferences.",
        ],
    )

    add_heading(doc, "13. Key Integrated Artifacts", 1)
    add_table(
        doc,
        ["Artifact", "Typical location", "Purpose in final platform"],
        [
            ["Processed cleaned dataset", "data/processed/cleaned_data.csv", "Provides the cleaned retail transaction foundation used for downstream analysis and dashboard summaries."],
            ["Feature-engineered dataset", "data/processed/features_data.csv and features_data_sample.csv", "Supports customer intelligence metrics, RFM-style summaries, and feature visibility for integration testing."],
            ["Feature summary", "dashboard/public/data/features_summary.json", "Lets the browser reflect the full engineered dataset without loading the full large CSV directly."],
            ["Demand forecasting outputs", "dashboard/public/data/demand_forecast.csv plus daily revenue and daily demand source files", "Drives forecast trend cards, confidence surfaces, and planning views."],
            ["Inventory outputs", "dashboard/public/data/inventory_recommendations.csv and SKU inventory source files", "Feeds reorder signals, stock posture, and inventory action tables."],
            ["Churn outputs", "dashboard/public/data/churn_predictions.csv and high_risk_customers.csv", "Powers risk distribution, revenue-at-risk, and retention targeting views."],
            ["Reporting and media layer", "dashboard/src/* and dashboard/public/data/*", "Generates board-ready previews, report pack narratives, and presentation-friendly dashboard states."],
        ],
        [1.45, 2.15, 2.6],
    )

    add_heading(doc, "14. Conclusion", 1)
    add_paragraph(doc, "The RetailPulse dashboard is a complete platform-integration product rather than a collection of isolated visual pages. It centralizes cleaned data, engineered features, forecasting outputs, churn analysis, inventory recommendations, and report-ready summaries into a coherent premium workspace. From both a technical and product-design perspective, it serves as the final experience layer that makes the broader RetailPulse project understandable, presentable, and operationally useful.")


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_document(doc)
    add_header_footer(doc)
    build_report(doc)
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
